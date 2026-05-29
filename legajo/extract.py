import re
import io
from typing import Optional

from legajo.dates import normalize_date, is_valid_date


# ============================================================
# PDF
# ============================================================


def extract_dni(text: str) -> Optional[str]:
    patterns = [
        r"DNI:\s*([\d\.]+)",
        r"D\.N\.I\.:\s*([\d\.]+)",
        r"DNI\s+([\d\.]+)",
        r"([\d]{1,2}\.[\d]{3}\.[\d]{3})",
    ]
    for pattern in patterns:
        m = re.search(pattern, text)
        if m:
            return m.group(1).replace(".", "")
    return None


def _normalizar_tipo(t: str):
    raw = t.upper()
    if raw in ("TIT.INTE", "TITULAR INTERINO", "INTERINO"):
        return "TITULAR"
    if raw.startswith("PROVISIO"):
        return "PROVISIONAL"
    return raw


def extract_periods_from_pdf(text: str):
    periods = []
    lines = [l.strip() for l in text.replace("\r\n", "\n").split("\n") if l.strip()]
    fecha_rx = re.compile(
        r"^(SUPLENTE|TIT(?:ULAR|\.INTE)|PROVISIO\w*|INTERINO)\s+"
        r"DEL\s+(\d{2}/\d{2}/\d{4})\s+AL\s+(\d{2}/\d{2}/\d{4})",
        re.IGNORECASE,
    )
    cargo_rx = re.compile(r"COMO\s+(.+?)\s*\d+\s*$", re.IGNORECASE)

    for i, line in enumerate(lines):
        m1 = fecha_rx.match(line)
        if not m1:
            continue
        tipo, fecha_inicio, fecha_fin = m1.group(1), m1.group(2), m1.group(3)
        cargo_source = line if "COMO" in line else ""
        if not cargo_source:
            for j in range(i + 1, min(i + 5, len(lines))):
                if len(lines[j].strip()) > 1 and "COMO" in lines[j]:
                    cargo_source = lines[j]
                    break
        m2 = cargo_rx.search(cargo_source)
        if not m2:
            continue
        cargo = re.sub(r"\s+", " ", m2.group(1).strip())
        d1, m1b, y1 = fecha_inicio.split("/")
        d2, m2b, y2 = fecha_fin.split("/")
        inicio = f"{y1}-{m1b}-{d1}"
        fin = f"{y2}-{m2b}-{d2}"
        if is_valid_date(inicio) and is_valid_date(fin) and inicio <= fin:
            periods.append({
                "inicio": inicio,
                "fin": fin,
                "situacion_revista": _normalizar_tipo(tipo),
            })
    return periods


def extract_licencias_from_pdf(text: str):
    licenses = []
    lines = [l.strip() for l in text.replace("\r\n", "\n").split("\n") if l.strip()]
    try:
        lic_idx = next(i for i, l in enumerate(lines) if l == "LICENCIAS")
    except StopIteration:
        return licenses

    fecha_rx = re.compile(r"(\d{2})/(\d{2})/(\d{4})\s*[-–]\s*(\d{2})/(\d{2})/(\d{4})")

    for i in range(lic_idx, len(lines)):
        m = fecha_rx.search(lines[i])
        if not m:
            continue
        d1, m1b, y1, d2, m2b, y2 = m.groups()
        inicio = f"{y1}-{m1b}-{d1}"
        fin = f"{y2}-{m2b}-{d2}"
        motivo = fecha_rx.sub("", lines[i]).strip()
        if not motivo:
            for j in range(i + 1, min(i + 5, len(lines))):
                nl = lines[j]
                if re.match(r"^\d{3}$", nl):
                    break
                if re.match(r"^_{3,}$", nl):
                    break
                motivo = nl
                break
        if is_valid_date(inicio) and is_valid_date(fin) and inicio <= fin:
            licenses.append({
                "inicio": inicio,
                "fin": fin,
                "motivo": motivo or "SIN ESPECIFICAR",
            })
    return licenses


# ============================================================
# EXCEL
# ============================================================


def _find_col_index(header, *names):
    lowered = [str(c).lower().strip() if c else "" for c in header]
    for name in names:
        nl = name.lower()
        for i, h in enumerate(lowered):
            if nl == h or nl in h or h in nl:
                return i
    return None


def _cell_to_str(cell):
    from datetime import datetime, date as dt_date

    if cell is None:
        return ""
    if isinstance(cell, (datetime, dt_date)):
        return cell.strftime("%Y-%m-%d")
    return str(cell)


def extract_periods_from_excel(buffer: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(buffer), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    periods = []
    if not rows:
        return periods
    header = list(rows[0]) if rows[0] else []

    inicio_idx = _find_col_index(header, "inicio", "fecha inicio", "fecha_inicio")
    fin_idx = _find_col_index(header, "fin", "fecha fin", "fecha_fin")
    if inicio_idx is None or fin_idx is None:
        return periods

    cargo_idx = _find_col_index(header, "cargo", "cargo")
    sit_idx = _find_col_index(header, "situacion_revista", "situacion", "tipo", "situacion revista")

    for row in rows[1:]:
        if not row or not row[inicio_idx] or not row[fin_idx]:
            continue
        inicio = normalize_date(_cell_to_str(row[inicio_idx]))
        fin = normalize_date(_cell_to_str(row[fin_idx]))
        if not inicio or not fin or not is_valid_date(inicio) or not is_valid_date(fin) or inicio > fin:
            continue
        cargo = _cell_to_str(row[cargo_idx]).strip() if cargo_idx is not None and row[cargo_idx] else "SIN_ESPECIFICAR"
        sit = _cell_to_str(row[sit_idx]).strip() if sit_idx is not None and row[sit_idx] else ""
        period = {"inicio": inicio, "fin": fin, "cargo": cargo}
        if sit:
            period["situacion_revista"] = sit.upper()
        periods.append(period)
    wb.close()
    return periods
