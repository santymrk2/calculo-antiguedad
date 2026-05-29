import asyncio
import io
import os
import re
import shutil

from fastapi import UploadFile, Request
from fastapi.responses import JSONResponse

from legajo.config import DATA_DIR
from legajo.storage import (
    ensure_dni_directory,
    read_legajo,
    write_legajo,
    read_licencias,
    write_licencias,
    read_config,
    write_config,
    save_source_document,
)
from legajo.extract import (
    extract_dni,
    extract_periods_from_pdf,
    extract_licencias_from_pdf,
    extract_periods_from_excel,
    _find_col_index,
    _cell_to_str,
)
from legajo.calc import (
    calculate_bruto,
    calculate_neto,
    calculate_by_role,
    create_normalized_record,
    is_duplicate,
)
from legajo.dates import days_to_commercial, days_to_natural


async def handle_process_file(request: Request):
    try:
        form = await request.form()
        file: UploadFile = form.get("file")
        if not file:
            return JSONResponse({"error": "No file provided"}, status_code=400)

        content = await file.read()
        filename = file.filename.lower()

        dni = None
        periods = []
        licencias = []

        if filename.endswith(".pdf"):
            import pdfplumber

            pdf = pdfplumber.open(io.BytesIO(content))
            text = ""
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
            pdf.close()

            dni = extract_dni(text)
            periods = extract_periods_from_pdf(text)
            licencias = extract_licencias_from_pdf(text)

        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            wb.close()

            if rows:
                header = list(rows[0]) if rows[0] else []
                dni_col = _find_col_index(header, "dni", "documento", "documento numero")
                if dni_col is not None and len(rows) > 1:
                    dni_raw = _cell_to_str(rows[1][dni_col])
                    if dni_raw:
                        dni = dni_raw.replace(".", "")

            periods = extract_periods_from_excel(content)

        else:
            return JSONResponse(
                {"error": "Format not supported (use PDF or Excel)"},
                status_code=400,
            )

        if not dni:
            return JSONResponse(
                {"error": "Could not extract DNI from file"},
                status_code=400,
            )

        if not periods and not licencias:
            return JSONResponse(
                {"error": "Could not extract periods from file"},
                status_code=400,
            )

        dni_dir, docs_dir, clean_dni = ensure_dni_directory(dni)
        safe_name, display_name = await save_source_document(docs_dir, content, file.filename)
        source_display = display_name

        legajo = await read_legajo(dni_dir)
        added_count = 0
        file_type = "PDF" if filename.endswith(".pdf") else "EXCEL"
        for period in periods:
            record = create_normalized_record(
                dni,
                period["inicio"],
                period["fin"],
                period.get("cargo", ""),
                f"{source_display} ({file_type})",
                period.get("situacion_revista", ""),
            )
            if not is_duplicate(record, legajo):
                legajo.append(record)
                added_count += 1

        await write_legajo(dni_dir, legajo)

        licencias_actuales = await read_licencias(dni_dir)
        for lic in licencias:
            if not any(l["inicio"] == lic["inicio"] and l["fin"] == lic["fin"] for l in licencias_actuales):
                licencias_actuales.append(lic)
        await write_licencias(dni_dir, licencias_actuales)

        config = await read_config(dni_dir)
        desc_lic = config.get("descontarLicencias", True) is not False

        comercial = {
            "bruto": calculate_bruto(legajo, days_to_commercial),
            "neto": calculate_neto(legajo, licencias_actuales, days_to_commercial, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_commercial),
        }
        natural = {
            "bruto": calculate_bruto(legajo, days_to_natural),
            "neto": calculate_neto(legajo, licencias_actuales, days_to_natural, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_natural),
        }

        return {
            "success": True,
            "dni": clean_dni,
            "recordsProcessed": len(periods),
            "recordsAdded": added_count,
            "totalRecords": len(legajo),
            "descontarLicencias": desc_lic,
            "comercial": comercial,
            "natural": natural,
            "legajo": legajo,
            "licencias": licencias_actuales,
        }

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_get_agente(dni_path):
    try:
        clean_dni = dni_path.replace(".", "")
        dni_dir = os.path.join(DATA_DIR, clean_dni)
        legajo = await read_legajo(dni_dir)

        if not legajo:
            return JSONResponse({"error": "Agent not found"}, status_code=404)

        licencias = await read_licencias(dni_dir)
        config = await read_config(dni_dir)
        desc_lic = config.get("descontarLicencias", True) is not False

        comercial = {
            "bruto": calculate_bruto(legajo, days_to_commercial),
            "neto": calculate_neto(legajo, licencias, days_to_commercial, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_commercial),
        }
        natural = {
            "bruto": calculate_bruto(legajo, days_to_natural),
            "neto": calculate_neto(legajo, licencias, days_to_natural, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_natural),
        }

        return {
            "dni": clean_dni,
            "totalRecords": len(legajo),
            "descontarLicencias": desc_lic,
            "comercial": comercial,
            "natural": natural,
            "legajo": legajo,
            "licencias": licencias,
        }

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_list_agentes():
    try:
        agentes = []
        for entry in os.scandir(DATA_DIR):
            if not entry.is_dir():
                continue
            dni = entry.name
            dni_dir = entry.path
            legajo = await read_legajo(dni_dir)
            if not legajo:
                continue

            docs_dir = os.path.join(dni_dir, "documentos_origen")
            doc_count = 0
            try:
                doc_count = len(os.listdir(docs_dir))
            except FileNotFoundError:
                pass

            licencias = await read_licencias(dni_dir)
            config = await read_config(dni_dir)
            desc_lic = config.get("descontarLicencias", True) is not False

            comercial = {
                "bruto": calculate_bruto(legajo, days_to_commercial),
                "neto": calculate_neto(legajo, licencias, days_to_commercial, desc_lic),
            }
            natural = {
                "bruto": calculate_bruto(legajo, days_to_natural),
                "neto": calculate_neto(legajo, licencias, days_to_natural, desc_lic),
            }

            primer_periodo = min(r["inicio"] for r in legajo)
            ultimo_periodo = max(r["fin"] for r in legajo)
            cargos = list(set(r.get("cargo", "") for r in legajo))

            agentes.append({
                "dni": dni,
                "totalRecords": len(legajo),
                "documentos": doc_count,
                "comercial": comercial,
                "natural": natural,
                "primerPeriodo": primer_periodo,
                "ultimoPeriodo": ultimo_periodo,
                "cargos": cargos,
            })

        agentes.sort(key=lambda a: a["dni"])
        return {"agentes": agentes}

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_list_documentos(dni):
    try:
        _, docs_dir, _ = ensure_dni_directory(dni)
        documentos = []
        for name in sorted(os.listdir(docs_dir)):
            m = re.match(r"^(\d+)_(.+)$", name)
            if m:
                documentos.append({
                    "name": name,
                    "displayName": m.group(2),
                    "uploadedAt": int(m.group(1)),
                })
            else:
                documentos.append({
                    "name": name,
                    "displayName": name,
                    "uploadedAt": None,
                })
        return {"documentos": documentos}
    except Exception:
        return {"documentos": []}


async def handle_reprocesar(dni):
    try:
        dni_dir, docs_dir, clean_dni = ensure_dni_directory(dni)
        try:
            files = os.listdir(docs_dir)
        except FileNotFoundError:
            return JSONResponse(
                {"error": "No hay documentos para reprocesar"},
                status_code=400,
            )

        if not files:
            return JSONResponse(
                {"error": "No hay documentos para reprocesar"},
                status_code=400,
            )

        import pdfplumber

        all_periods = []
        todas_licencias = []

        for file in files:
            file_path = os.path.join(docs_dir, file)
            loop = asyncio.get_event_loop()

            def _read():
                with open(file_path, "rb") as f:
                    return f.read()

            buffer = await loop.run_in_executor(None, _read)
            fname = file.lower()

            if fname.endswith(".pdf"):
                pdf = pdfplumber.open(io.BytesIO(buffer))
                text = ""
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text += t + "\n"
                pdf.close()

                periods = extract_periods_from_pdf(text)
                licencias = extract_licencias_from_pdf(text)
                for lic in licencias:
                    if not any(l["inicio"] == lic["inicio"] and l["fin"] == lic["fin"] for l in todas_licencias):
                        todas_licencias.append(lic)
            elif fname.endswith(".xlsx") or fname.endswith(".xls"):
                periods = extract_periods_from_excel(buffer)
            else:
                continue

            display_name = re.sub(r"^\d+_", "", file)
            file_type = "PDF" if fname.endswith(".pdf") else "EXCEL"
            for p in periods:
                p["origen"] = f"{display_name} ({file_type})"
                all_periods.append(p)

        legajo = []
        added_count = 0
        for period in all_periods:
            record = create_normalized_record(
                clean_dni,
                period["inicio"],
                period["fin"],
                period.get("cargo", ""),
                period.get("origen", ""),
                period.get("situacion_revista", ""),
            )
            if not is_duplicate(record, legajo):
                legajo.append(record)
                added_count += 1

        await write_legajo(dni_dir, legajo)
        await write_licencias(dni_dir, todas_licencias)

        config = await read_config(dni_dir)
        desc_lic = config.get("descontarLicencias", True) is not False

        comercial = {
            "bruto": calculate_bruto(legajo, days_to_commercial),
            "neto": calculate_neto(legajo, todas_licencias, days_to_commercial, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_commercial),
        }
        natural = {
            "bruto": calculate_bruto(legajo, days_to_natural),
            "neto": calculate_neto(legajo, todas_licencias, days_to_natural, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_natural),
        }

        return {
            "success": True,
            "dni": clean_dni,
            "recordsProcessed": len(all_periods),
            "recordsAdded": added_count,
            "totalRecords": len(legajo),
            "descontarLicencias": desc_lic,
            "comercial": comercial,
            "natural": natural,
            "legajo": legajo,
            "licencias": todas_licencias,
        }

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_delete_documento(dni, filename):
    try:
        dni_dir, docs_dir, clean_dni = ensure_dni_directory(dni)
        file_path = os.path.join(docs_dir, filename)

        if not os.path.exists(file_path):
            return JSONResponse(
                {"error": "El archivo no existe"},
                status_code=404,
            )

        os.remove(file_path)

        legajo = await read_legajo(dni_dir)
        before_count = len(legajo)
        legajo = [r for r in legajo if not r.get("origen", "").startswith(filename)]
        removed_count = before_count - len(legajo)
        await write_legajo(dni_dir, legajo)

        response_data = {
            "success": True,
            "dni": clean_dni,
            "archivoEliminado": filename,
            "registrosRemovidos": removed_count,
            "totalRecords": len(legajo),
        }

        if legajo:
            licencias = await read_licencias(dni_dir)
            config = await read_config(dni_dir)
            desc_lic = config.get("descontarLicencias", True) is not False
            response_data["comercial"] = {
                "bruto": calculate_bruto(legajo, days_to_commercial),
                "neto": calculate_neto(legajo, licencias, days_to_commercial, desc_lic),
                "byRole": calculate_by_role(legajo, days_to_commercial),
            }
            response_data["natural"] = {
                "bruto": calculate_bruto(legajo, days_to_natural),
                "neto": calculate_neto(legajo, licencias, days_to_natural, desc_lic),
                "byRole": calculate_by_role(legajo, days_to_natural),
            }
            response_data["legajo"] = legajo
            response_data["licencias"] = licencias
            response_data["descontarLicencias"] = desc_lic

        return response_data

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_clear_agente(dni):
    try:
        clean_dni = dni.replace(".", "")
        dni_dir = os.path.join(DATA_DIR, clean_dni)

        if not os.path.exists(dni_dir):
            return JSONResponse(
                {"error": "El agente no existe"},
                status_code=404,
            )

        shutil.rmtree(dni_dir)
        return {"success": True, "dni": clean_dni, "eliminado": True}

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_update_config(dni, request):
    try:
        clean_dni = dni.replace(".", "")
        dni_dir = os.path.join(DATA_DIR, clean_dni)
        body = await request.json()
        config = await write_config(dni_dir, body)

        legajo = await read_legajo(dni_dir)
        licencias = await read_licencias(dni_dir)
        desc_lic = config.get("descontarLicencias", True) is not False

        comercial = {
            "bruto": calculate_bruto(legajo, days_to_commercial),
            "neto": calculate_neto(legajo, licencias, days_to_commercial, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_commercial),
        }
        natural = {
            "bruto": calculate_bruto(legajo, days_to_natural),
            "neto": calculate_neto(legajo, licencias, days_to_natural, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_natural),
        }

        return {
            "success": True,
            "dni": clean_dni,
            "descontarLicencias": desc_lic,
            "comercial": comercial,
            "natural": natural,
            "legajo": legajo,
            "licencias": licencias,
        }

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)
