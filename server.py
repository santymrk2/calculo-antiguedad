import asyncio
import json
import math
import os
import re
import shutil
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

from fastapi import FastAPI, File, Form, HTTPException, UploadFile, Request
from fastapi.responses import HTMLResponse, JSONResponse
import uvicorn

app = FastAPI(title="Legajo Digital")

DATA_DIR = os.environ.get(
    "DATA_DIR",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "datos_locales"),
)
PORT = int(os.environ.get("PORT", "3000"))
HOST = os.environ.get("HOST", "0.0.0.0")

# ============================================================================
# UTILIDADES DE SISTEMA DE ARCHIVOS
# ============================================================================


def ensure_dni_directory(dni: str):
    clean_dni = dni.replace(".", "")
    dni_dir = os.path.join(DATA_DIR, clean_dni)
    docs_dir = os.path.join(dni_dir, "documentos_origen")
    os.makedirs(docs_dir, exist_ok=True)
    return dni_dir, docs_dir, clean_dni


async def read_legajo(dni_dir: str):
    path = os.path.join(dni_dir, "legajo_historico.json")
    try:
        loop = asyncio.get_event_loop()
        with open(path, "r") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return []


async def write_legajo(dni_dir: str, legajo: list):
    path = os.path.join(dni_dir, "legajo_historico.json")
    loop = asyncio.get_event_loop()

    def _write():
        with open(path, "w") as f:
            json.dump(legajo, f, indent=2, ensure_ascii=False)

    await loop.run_in_executor(None, _write)
    return True


async def read_licencias(dni_dir: str):
    path = os.path.join(dni_dir, "licencias.json")
    try:
        loop = asyncio.get_event_loop()
        with open(path, "r") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return []


async def write_licencias(dni_dir: str, licencias: list):
    path = os.path.join(dni_dir, "licencias.json")
    loop = asyncio.get_event_loop()

    def _write():
        with open(path, "w") as f:
            json.dump(licencias, f, indent=2, ensure_ascii=False)

    await loop.run_in_executor(None, _write)
    return True


async def read_config(dni_dir: str):
    path = os.path.join(dni_dir, "config.json")
    try:
        loop = asyncio.get_event_loop()
        with open(path, "r") as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


async def write_config(dni_dir: str, updates: dict):
    current = await read_config(dni_dir)
    current.update(updates)
    path = os.path.join(dni_dir, "config.json")
    loop = asyncio.get_event_loop()

    def _write():
        with open(path, "w") as f:
            json.dump(current, f, indent=2, ensure_ascii=False)

    await loop.run_in_executor(None, _write)
    return current


async def save_source_document(docs_dir: str, content: bytes, filename: str):
    timestamp = int(datetime.now().timestamp() * 1000)
    clean_name = re.sub(r"^[^a-zA-Z0-9]+", "", filename)
    safe_name = f"{timestamp}_{clean_name}"
    file_path = os.path.join(docs_dir, safe_name)

    loop = asyncio.get_event_loop()

    def _write():
        with open(file_path, "wb") as f:
            f.write(content)

    await loop.run_in_executor(None, _write)
    return safe_name, clean_name


# ============================================================================
# EXTRACCIÓN DE DATOS (PDF)
# ============================================================================


def extract_dni(text: str) -> Optional[str]:
    patterns = [
        r"DNI:\s*([\d\.]+)",
        r"D\.N\.I\.:\s*([\d\.]+)",
        r"DNI\s+([\d\.]+)",
        r"([\d]{1,2}\.[\d]{3}\.[\d]{3})",
    ]
    for pattern in patterns:
        m = re.search(pattern, text)
        if m:
            return m.group(1).replace(".", "")
    return None


def extract_periods_from_pdf(text: str):
    periods = []
    lines = [
        l.strip() for l in text.replace("\r\n", "\n").split("\n") if l.strip()
    ]
    fecha_rx = re.compile(
        r"^(SUPLENTE|TIT(?:ULAR|\.INTE)|PROVISIO\w*|INTERINO)\s+"
        r"DEL\s+(\d{2}/\d{2}/\d{4})\s+AL\s+(\d{2}/\d{2}/\d{4})",
        re.IGNORECASE,
    )
    cargo_rx = re.compile(r"COMO\s+(.+?)\s*\d+\s*$", re.IGNORECASE)

    def normalizar_tipo(t: str):
        raw = t.upper()
        if raw in ("TIT.INTE", "TITULAR INTERINO", "INTERINO"):
            return "TITULAR"
        if raw.startswith("PROVISIO"):
            return "PROVISIONAL"
        return raw

    for i, line in enumerate(lines):
        m1 = fecha_rx.match(line)
        if not m1:
            continue
        tipo, fecha_inicio, fecha_fin = m1.group(1), m1.group(2), m1.group(3)
        cargo_source = line if "COMO" in line else (lines[i + 1] if i + 1 < len(lines) else "")
        m2 = cargo_rx.search(cargo_source)
        if not m2:
            continue
        cargo = re.sub(r"\s+", " ", m2.group(1).strip())
        d1, m1b, y1 = fecha_inicio.split("/")
        d2, m2b, y2 = fecha_fin.split("/")
        inicio = f"{y1}-{m1b}-{d1}"
        fin = f"{y2}-{m2b}-{d2}"
        if is_valid_date(inicio) and is_valid_date(fin) and inicio <= fin:
            periods.append({
                "inicio": inicio,
                "fin": fin,
                "situacion_revista": normalizar_tipo(tipo),
            })
    return periods


def extract_licencias_from_pdf(text: str):
    licenses = []
    lines = [
        l.strip() for l in text.replace("\r\n", "\n").split("\n") if l.strip()
    ]
    try:
        lic_idx = next(i for i, l in enumerate(lines) if l == "LICENCIAS")
    except StopIteration:
        return licenses

    fecha_rx = re.compile(
        r"(\d{2})/(\d{2})/(\d{4})\s*[-–]\s*(\d{2})/(\d{2})/(\d{4})"
    )

    for i in range(lic_idx, len(lines)):
        m = fecha_rx.search(lines[i])
        if not m:
            continue
        d1, m1b, y1, d2, m2b, y2 = m.groups()
        inicio = f"{y1}-{m1b}-{d1}"
        fin = f"{y2}-{m2b}-{d2}"
        motivo = fecha_rx.sub("", lines[i]).strip()
        if not motivo:
            for j in range(i + 1, min(i + 5, len(lines))):
                nl = lines[j]
                if re.match(r"^\d{3}$", nl):
                    break
                if re.match(r"^_{3,}$", nl):
                    break
                motivo = nl
                break
        if is_valid_date(inicio) and is_valid_date(fin) and inicio <= fin:
            licenses.append({
                "inicio": inicio,
                "fin": fin,
                "motivo": motivo or "SIN ESPECIFICAR",
            })
    return licenses


# ============================================================================
# EXTRACCIÓN DE DATOS (EXCEL)
# ============================================================================


def _find_col_index(header, *names):
    lowered = [str(c).lower().strip() if c else "" for c in header]
    for name in names:
        nl = name.lower()
        for i, h in enumerate(lowered):
            if nl == h or nl in h or h in nl:
                return i
    return None


def _cell_to_str(cell):
    from datetime import datetime, date as dt_date

    if cell is None:
        return ""
    if isinstance(cell, (datetime, dt_date)):
        return cell.strftime("%Y-%m-%d")
    return str(cell)


def extract_periods_from_excel(buffer: bytes):
    from openpyxl import load_workbook
    import io

    wb = load_workbook(io.BytesIO(buffer), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    periods = []
    if not rows:
        return periods
    header = list(rows[0]) if rows[0] else []

    inicio_idx = _find_col_index(header, "inicio", "fecha inicio", "fecha_inicio")
    fin_idx = _find_col_index(header, "fin", "fecha fin", "fecha_fin")
    if inicio_idx is None or fin_idx is None:
        return periods

    cargo_idx = _find_col_index(header, "cargo", "cargo")
    sit_idx = _find_col_index(header, "situacion_revista", "situacion", "tipo", "situacion revista")

    for row in rows[1:]:
        if not row or not row[inicio_idx] or not row[fin_idx]:
            continue
        inicio = normalize_date(_cell_to_str(row[inicio_idx]))
        fin = normalize_date(_cell_to_str(row[fin_idx]))
        if not inicio or not fin or not is_valid_date(inicio) or not is_valid_date(fin) or inicio > fin:
            continue
        cargo = _cell_to_str(row[cargo_idx]).strip() if cargo_idx is not None and row[cargo_idx] else "SIN_ESPECIFICAR"
        sit = _cell_to_str(row[sit_idx]).strip() if sit_idx is not None and row[sit_idx] else ""
        period = {"inicio": inicio, "fin": fin, "cargo": cargo}
        if sit:
            period["situacion_revista"] = sit.upper()
        periods.append(period)
    wb.close()
    return periods


# ============================================================================
# UTILIDADES DE FECHAS
# ============================================================================


def normalize_date(date_str: str) -> Optional[str]:
    from datetime import datetime, date as dt_date

    if not date_str:
        return None
    if isinstance(date_str, (datetime, dt_date)):
        return date_str.strftime("%Y-%m-%d")
    s = date_str.strip()
    # YYYY-MM-DD
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        return s
    # DD/MM/YYYY
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    # DD-MM-YYYY
    m = re.match(r"^(\d{1,2})-(\d{1,2})-(\d{4})$", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    # YYYY/MM/DD
    m = re.match(r"^(\d{4})/(\d{1,2})/(\d{1,2})$", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    # Excel serial number
    if re.match(r"^\d{4,5}$", s):
        from datetime import datetime, timedelta

        serial = int(s)
        dt = datetime(1899, 12, 30) + timedelta(days=serial)
        return dt.strftime("%Y-%m-%d")
    return None


def is_valid_date(date_str: str) -> bool:
    if not date_str or not re.match(r"^\d{4}-\d{2}-\d{2}$", date_str):
        return False
    try:
        y, m, d = map(int, date_str.split("-"))
        dt = date(y, m, d)
        return dt.year == y and dt.month == m and dt.day == d
    except ValueError:
        return False


def days_between_natural(inicio: str, fin: str) -> int:
    start = date.fromisoformat(inicio)
    end = date.fromisoformat(fin)
    diff = abs((end - start).days)
    return diff + 1  # inclusive


def days_to_commercial(natural_days: int):
    years = natural_days // 360
    remainder = natural_days % 360
    months = remainder // 30
    days = remainder % 30
    return {"years": years, "months": months, "days": days, "total": natural_days}


def days_to_natural(natural_days: int):
    years = natural_days // 365
    remainder = natural_days % 365
    months = remainder // 30
    days = remainder % 30
    return {"years": years, "months": months, "days": days, "total": natural_days}


# ============================================================================
# CÁLCULO DE ANTIGÜEDAD
# ============================================================================


def merge_intervals(periods: list):
    if not periods:
        return []
    sorted_p = sorted(periods, key=lambda p: p["inicio"])
    merged = [dict(sorted_p[0])]
    for p in sorted_p[1:]:
        last = merged[-1]
        if p["inicio"] <= last["fin"]:
            if p["fin"] > last["fin"]:
                last["fin"] = p["fin"]
        else:
            merged.append(dict(p))
    return merged


def add_days(date_str: str, days: int):
    dt = date.fromisoformat(date_str)
    dt += timedelta(days=days)
    return dt.isoformat()


def subtract_intervals(intervals, subtractions):
    result = []
    for interval in intervals:
        parts = [{"inicio": interval["inicio"], "fin": interval["fin"]}]
        for sub in subtractions:
            new_parts = []
            for part in parts:
                if sub["fin"] < part["inicio"] or sub["inicio"] > part["fin"]:
                    new_parts.append(part)
                    continue
                if sub["inicio"] > part["inicio"]:
                    new_parts.append({"inicio": part["inicio"], "fin": add_days(sub["inicio"], -1)})
                if sub["fin"] < part["fin"]:
                    new_parts.append({"inicio": add_days(sub["fin"], 1), "fin": part["fin"]})
            parts = new_parts
        result.extend(parts)
    return result


def calculate_bruto(periods, converter=days_to_commercial):
    total = sum(days_between_natural(p["inicio"], p["fin"]) for p in periods)
    return converter(total)


def calculate_neto(periods, licencias=None, converter=days_to_commercial, descontar_licencias=True):
    if licencias is None:
        licencias = []
    merged = merge_intervals(periods)
    intervals = subtract_intervals(merged, licencias) if descontar_licencias else merged
    return calculate_bruto(intervals, converter)


def calculate_by_role(periods, converter=days_to_commercial):
    by_role = {}
    for p in periods:
        key = p.get("situacion_revista") or "SIN ESPECIFICAR"
        if key not in by_role:
            by_role[key] = []
        by_role[key].append(p)
    return {role: calculate_bruto(role_periods, converter) for role, role_periods in by_role.items()}


def create_normalized_record(dni: str, inicio: str, fin: str, cargo: str, origen: str, situacion_revista: str):
    return {
        "dni": dni.replace(".", ""),
        "inicio": inicio,
        "fin": fin,
        "cargo": cargo.strip(),
        "situacion_revista": (situacion_revista or "").upper(),
        "origen": origen,
    }


def is_duplicate(record: dict, legajo: list):
    return any(
        item["dni"] == record["dni"]
        and item["inicio"] == record["inicio"]
        and item["fin"] == record["fin"]
        and item["cargo"] == record["cargo"]
        and item["situacion_revista"] == record["situacion_revista"]
        for item in legajo
    )


# ============================================================================
# HANDLERS
# ============================================================================


async def handle_process_file(request: Request):
    import io

    try:
        form = await request.form()
        file: UploadFile = form.get("file")
        if not file:
            return JSONResponse({"error": "No file provided"}, status_code=400)

        content = await file.read()
        filename = file.filename.lower()

        dni = None
        periods = []
        licencias = []

        if filename.endswith(".pdf"):
            import pdfplumber

            pdf = pdfplumber.open(io.BytesIO(content))
            text = ""
            for page in pdf.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
            pdf.close()

            dni = extract_dni(text)
            periods = extract_periods_from_pdf(text)
            licencias = extract_licencias_from_pdf(text)

        elif filename.endswith(".xlsx") or filename.endswith(".xls"):
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            wb.close()

            if rows:
                header = list(rows[0]) if rows[0] else []
                dni_col = _find_col_index(header, "dni", "documento", "documento numero")
                if dni_col is not None and len(rows) > 1:
                    dni_raw = _cell_to_str(rows[1][dni_col])
                    if dni_raw:
                        dni = dni_raw.replace(".", "")

            periods = extract_periods_from_excel(content)

        else:
            return JSONResponse(
                {"error": "Format not supported (use PDF or Excel)"},
                status_code=400,
            )

        if not dni:
            return JSONResponse(
                {"error": "Could not extract DNI from file"},
                status_code=400,
            )

        if not periods and not licencias:
            return JSONResponse(
                {"error": "Could not extract periods from file"},
                status_code=400,
            )

        dni_dir, docs_dir, clean_dni = ensure_dni_directory(dni)
        safe_name, display_name = await save_source_document(docs_dir, content, file.filename)
        source_display = display_name

        legajo = await read_legajo(dni_dir)
        added_count = 0
        file_type = "PDF" if filename.endswith(".pdf") else "EXCEL"
        for period in periods:
            record = create_normalized_record(
                dni,
                period["inicio"],
                period["fin"],
                period.get("cargo", ""),
                f"{source_display} ({file_type})",
                period.get("situacion_revista", ""),
            )
            if not is_duplicate(record, legajo):
                legajo.append(record)
                added_count += 1

        await write_legajo(dni_dir, legajo)

        licencias_actuales = await read_licencias(dni_dir)
        for lic in licencias:
            if not any(l["inicio"] == lic["inicio"] and l["fin"] == lic["fin"] for l in licencias_actuales):
                licencias_actuales.append(lic)
        await write_licencias(dni_dir, licencias_actuales)

        config = await read_config(dni_dir)
        desc_lic = config.get("descontarLicencias", True) is not False

        comercial = {
            "bruto": calculate_bruto(legajo, days_to_commercial),
            "neto": calculate_neto(legajo, licencias_actuales, days_to_commercial, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_commercial),
        }
        natural = {
            "bruto": calculate_bruto(legajo, days_to_natural),
            "neto": calculate_neto(legajo, licencias_actuales, days_to_natural, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_natural),
        }

        return {
            "success": True,
            "dni": clean_dni,
            "recordsProcessed": len(periods),
            "recordsAdded": added_count,
            "totalRecords": len(legajo),
            "descontarLicencias": desc_lic,
            "comercial": comercial,
            "natural": natural,
            "legajo": legajo,
            "licencias": licencias_actuales,
        }

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_get_agente(dni_path: str):
    try:
        clean_dni = dni_path.replace(".", "")
        dni_dir = os.path.join(DATA_DIR, clean_dni)
        legajo = await read_legajo(dni_dir)

        if not legajo:
            return JSONResponse({"error": "Agent not found"}, status_code=404)

        licencias = await read_licencias(dni_dir)
        config = await read_config(dni_dir)
        desc_lic = config.get("descontarLicencias", True) is not False

        comercial = {
            "bruto": calculate_bruto(legajo, days_to_commercial),
            "neto": calculate_neto(legajo, licencias, days_to_commercial, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_commercial),
        }
        natural = {
            "bruto": calculate_bruto(legajo, days_to_natural),
            "neto": calculate_neto(legajo, licencias, days_to_natural, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_natural),
        }

        return {
            "dni": clean_dni,
            "totalRecords": len(legajo),
            "descontarLicencias": desc_lic,
            "comercial": comercial,
            "natural": natural,
            "legajo": legajo,
            "licencias": licencias,
        }

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_list_agentes():
    try:
        agentes = []
        for entry in os.scandir(DATA_DIR):
            if not entry.is_dir():
                continue
            dni = entry.name
            dni_dir = entry.path
            legajo = await read_legajo(dni_dir)
            if not legajo:
                continue

            docs_dir = os.path.join(dni_dir, "documentos_origen")
            doc_count = 0
            try:
                doc_count = len(os.listdir(docs_dir))
            except FileNotFoundError:
                pass

            licencias = await read_licencias(dni_dir)
            config = await read_config(dni_dir)
            desc_lic = config.get("descontarLicencias", True) is not False

            comercial = {
                "bruto": calculate_bruto(legajo, days_to_commercial),
                "neto": calculate_neto(legajo, licencias, days_to_commercial, desc_lic),
            }
            natural = {
                "bruto": calculate_bruto(legajo, days_to_natural),
                "neto": calculate_neto(legajo, licencias, days_to_natural, desc_lic),
            }

            primer_periodo = min(r["inicio"] for r in legajo)
            ultimo_periodo = max(r["fin"] for r in legajo)
            cargos = list(set(r.get("cargo", "") for r in legajo))

            agentes.append({
                "dni": dni,
                "totalRecords": len(legajo),
                "documentos": doc_count,
                "comercial": comercial,
                "natural": natural,
                "primerPeriodo": primer_periodo,
                "ultimoPeriodo": ultimo_periodo,
                "cargos": cargos,
            })

        agentes.sort(key=lambda a: a["dni"])
        return {"agentes": agentes}

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_list_documentos(dni: str):
    try:
        _, docs_dir, _ = ensure_dni_directory(dni)
        documentos = []
        for name in sorted(os.listdir(docs_dir)):
            m = re.match(r"^(\d+)_(.+)$", name)
            if m:
                documentos.append({
                    "name": name,
                    "displayName": m.group(2),
                    "uploadedAt": int(m.group(1)),
                })
            else:
                documentos.append({
                    "name": name,
                    "displayName": name,
                    "uploadedAt": None,
                })
        return {"documentos": documentos}
    except Exception:
        return {"documentos": []}


async def handle_reprocesar(dni: str):
    try:
        dni_dir, docs_dir, clean_dni = ensure_dni_directory(dni)
        files = []
        try:
            files = os.listdir(docs_dir)
        except FileNotFoundError:
            return JSONResponse(
                {"error": "No hay documentos para reprocesar"},
                status_code=400,
            )

        if not files:
            return JSONResponse(
                {"error": "No hay documentos para reprocesar"},
                status_code=400,
            )

        import pdfplumber
        import io

        all_periods = []
        todas_licencias = []

        for file in files:
            file_path = os.path.join(docs_dir, file)
            loop = asyncio.get_event_loop()

            def _read():
                with open(file_path, "rb") as f:
                    return f.read()

            buffer = await loop.run_in_executor(None, _read)
            fname = file.lower()

            if fname.endswith(".pdf"):
                pdf = pdfplumber.open(io.BytesIO(buffer))
                text = ""
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        text += t + "\n"
                pdf.close()

                periods = extract_periods_from_pdf(text)
                licencias = extract_licencias_from_pdf(text)
                for lic in licencias:
                    if not any(l["inicio"] == lic["inicio"] and l["fin"] == lic["fin"] for l in todas_licencias):
                        todas_licencias.append(lic)
            elif fname.endswith(".xlsx") or fname.endswith(".xls"):
                periods = extract_periods_from_excel(buffer)
            else:
                continue

            display_name = re.sub(r"^\d+_", "", file)
            file_type = "PDF" if fname.endswith(".pdf") else "EXCEL"
            for p in periods:
                p["origen"] = f"{display_name} ({file_type})"
                all_periods.append(p)

        legajo = []
        added_count = 0
        for period in all_periods:
            record = create_normalized_record(
                clean_dni,
                period["inicio"],
                period["fin"],
                period.get("cargo", ""),
                period.get("origen", ""),
                period.get("situacion_revista", ""),
            )
            if not is_duplicate(record, legajo):
                legajo.append(record)
                added_count += 1

        await write_legajo(dni_dir, legajo)
        await write_licencias(dni_dir, todas_licencias)

        config = await read_config(dni_dir)
        desc_lic = config.get("descontarLicencias", True) is not False

        comercial = {
            "bruto": calculate_bruto(legajo, days_to_commercial),
            "neto": calculate_neto(legajo, todas_licencias, days_to_commercial, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_commercial),
        }
        natural = {
            "bruto": calculate_bruto(legajo, days_to_natural),
            "neto": calculate_neto(legajo, todas_licencias, days_to_natural, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_natural),
        }

        return {
            "success": True,
            "dni": clean_dni,
            "recordsProcessed": len(all_periods),
            "recordsAdded": added_count,
            "totalRecords": len(legajo),
            "descontarLicencias": desc_lic,
            "comercial": comercial,
            "natural": natural,
            "legajo": legajo,
            "licencias": todas_licencias,
        }

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_delete_documento(dni: str, filename: str):
    try:
        dni_dir, docs_dir, clean_dni = ensure_dni_directory(dni)
        file_path = os.path.join(docs_dir, filename)

        if not os.path.exists(file_path):
            return JSONResponse(
                {"error": "El archivo no existe"},
                status_code=404,
            )

        os.remove(file_path)

        legajo = await read_legajo(dni_dir)
        before_count = len(legajo)
        legajo = [r for r in legajo if not r.get("origen", "").startswith(filename)]
        removed_count = before_count - len(legajo)
        await write_legajo(dni_dir, legajo)

        response_data = {
            "success": True,
            "dni": clean_dni,
            "archivoEliminado": filename,
            "registrosRemovidos": removed_count,
            "totalRecords": len(legajo),
        }

        if legajo:
            licencias = await read_licencias(dni_dir)
            config = await read_config(dni_dir)
            desc_lic = config.get("descontarLicencias", True) is not False
            response_data["comercial"] = {
                "bruto": calculate_bruto(legajo, days_to_commercial),
                "neto": calculate_neto(legajo, licencias, days_to_commercial, desc_lic),
                "byRole": calculate_by_role(legajo, days_to_commercial),
            }
            response_data["natural"] = {
                "bruto": calculate_bruto(legajo, days_to_natural),
                "neto": calculate_neto(legajo, licencias, days_to_natural, desc_lic),
                "byRole": calculate_by_role(legajo, days_to_natural),
            }
            response_data["legajo"] = legajo
            response_data["licencias"] = licencias
            response_data["descontarLicencias"] = desc_lic

        return response_data

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_clear_agente(dni: str):
    try:
        clean_dni = dni.replace(".", "")
        dni_dir = os.path.join(DATA_DIR, clean_dni)

        if not os.path.exists(dni_dir):
            return JSONResponse(
                {"error": "El agente no existe"},
                status_code=404,
            )

        shutil.rmtree(dni_dir)
        return {"success": True, "dni": clean_dni, "eliminado": True}

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


async def handle_update_config(dni: str, request: Request):
    try:
        clean_dni = dni.replace(".", "")
        dni_dir = os.path.join(DATA_DIR, clean_dni)
        body = await request.json()
        config = await write_config(dni_dir, body)

        legajo = await read_legajo(dni_dir)
        licencias = await read_licencias(dni_dir)
        desc_lic = config.get("descontarLicencias", True) is not False

        comercial = {
            "bruto": calculate_bruto(legajo, days_to_commercial),
            "neto": calculate_neto(legajo, licencias, days_to_commercial, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_commercial),
        }
        natural = {
            "bruto": calculate_bruto(legajo, days_to_natural),
            "neto": calculate_neto(legajo, licencias, days_to_natural, desc_lic),
            "byRole": calculate_by_role(legajo, days_to_natural),
        }

        return {
            "success": True,
            "dni": clean_dni,
            "descontarLicencias": desc_lic,
            "comercial": comercial,
            "natural": natural,
            "legajo": legajo,
            "licencias": licencias,
        }

    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


# ============================================================================
# ROUTES
# ============================================================================


@app.get("/")
async def root():
    html_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "index.html")
    try:
        with open(html_path, "r") as f:
            return HTMLResponse(f.read())
    except FileNotFoundError:
        return HTMLResponse("index.html not found", status_code=404)


@app.post("/procesar")
async def procesar(request: Request):
    return await handle_process_file(request)


@app.get("/agentes")
async def list_agentes():
    return await handle_list_agentes()


@app.get("/agente/{dni}/documentos")
async def list_documentos(dni: str):
    return await handle_list_documentos(dni)


@app.get("/agente/{dni}")
async def get_agente(dni: str):
    return await handle_get_agente(dni)


@app.delete("/agente/{dni}")
async def clear_agente(dni: str):
    return await handle_clear_agente(dni)


@app.post("/reprocesar/{dni}")
async def reprocesar(dni: str):
    return await handle_reprocesar(dni)


@app.put("/agente/{dni}/config")
async def update_config(dni: str, request: Request):
    return await handle_update_config(dni, request)


@app.delete("/documento/{dni:path}")
async def delete_documento(dni: str, request: Request):
    path_parts = dni.split("/")
    if len(path_parts) >= 2:
        dni_val = path_parts[0]
        filename = "/".join(path_parts[1:])
        from urllib.parse import unquote

        filename = unquote(filename)
        return await handle_delete_documento(dni_val, filename)
    return JSONResponse({"error": "Invalid path"}, status_code=400)


# ============================================================================
# MAIN
# ============================================================================

if __name__ == "__main__":
    os.makedirs(DATA_DIR, exist_ok=True)
    print(f"✅ Servidor iniciado en http://localhost:{PORT}")
    print(f"📁 Datos almacenados en: {DATA_DIR}")
    uvicorn.run(app, host=HOST, port=PORT)
